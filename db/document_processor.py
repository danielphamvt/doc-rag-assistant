import logging
import os
import sys

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import glob
import json
from pathlib import Path

import pymupdf
import pymupdf4llm
from langchain_qdrant import QdrantVectorStore
from langchain_qdrant.qdrant import RetrievalMode
from langchain_text_splitters import (
    MarkdownHeaderTextSplitter,
    RecursiveCharacterTextSplitter,
)
from qdrant_client.http import models as qmodels

logger = logging.getLogger(__name__)

from config import (
    CHILD_CHUNK_OVERLAP,
    CHILD_CHUNK_SIZE,
    CHILD_COLLECTION,
    DOCS_DIR,
    INDEX_BATCH_SIZE,
    MARKDOWN_DIR,
    MAX_PARENT_SIZE,
    MIN_PARENT_SIZE,
    PARENT_STORE_PATH,
    client,
    dense_embeddings,
    sparse_embeddings,
)
from domain_profile import PROFILE

os.environ["TOKENIZERS_PARALLELISM"] = "false"

# --- 1. PDF to Markdown Conversion ---
def pdf_to_markdown(pdf_path: Path, output_dir: Path):
    doc = pymupdf.open(pdf_path)
    md = pymupdf4llm.to_markdown(
        doc, 
        header=False, 
        footer=False, 
        page_separators=True, 
        ignore_images=True, 
        write_images=False, 
        image_path=None
    )
    # Clean unicode errors
    md_cleaned = md.encode('utf-8', errors='surrogatepass').decode('utf-8', errors='ignore')
    output_path = output_dir / pdf_path.stem
    output_path.with_suffix(".md").write_bytes(md_cleaned.encode('utf-8'))

def docx_to_markdown(docx_path: Path, output_dir: Path):
    import docx
    try:
        doc = docx.Document(docx_path)
    except Exception as e:
        logger.error("Failed to open docx file %s: %s", docx_path.name, e)
        return

    def paragraph_to_md(p) -> str:
        text = p.text.strip()
        if not text:
            return ""
        style = p.style.name.lower()
        if style.startswith("heading 1"):
            return f"\n# {text}\n"
        if style.startswith("heading 2"):
            return f"\n## {text}\n"
        if style.startswith("heading 3"):
            return f"\n### {text}\n"
        return text

    # Map element -> paragraph/table so the body can be walked in original order (O(1) per element)
    paragraphs_by_el = {p._element: p for p in doc.paragraphs}
    tables_by_el = {t._element: t for t in doc.tables}

    md_lines = []
    try:
        for element in doc.element.body:
            if element in paragraphs_by_el:
                md_lines.append(paragraph_to_md(paragraphs_by_el[element]))
            elif element in tables_by_el:
                md_lines.append("")
                for row in tables_by_el[element].rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    md_lines.append("| " + " | ".join(row_cells) + " |")
                md_lines.append("")
    except Exception as e:
        logger.warning("Error during ordered traversal of docx %s: %s. Falling back to paragraphs only.", docx_path.name, e)
        md_lines = []

    # Fallback to simple paragraph extraction if body element traversal failed or found nothing
    if not md_lines:
        md_lines = [paragraph_to_md(p) for p in doc.paragraphs]

    md_content = "\n".join(md_lines)
    md_cleaned = md_content.encode('utf-8', errors='surrogatepass').decode('utf-8', errors='ignore')
    output_path = output_dir / docx_path.stem
    output_path.with_suffix(".md").write_bytes(md_cleaned.encode('utf-8'))

def calculate_md5(file_path: Path) -> str:
    import hashlib
    hash_md5 = hashlib.md5()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            hash_md5.update(chunk)
    return hash_md5.hexdigest()

def documents_to_markdowns(changed_files: list):
    output_dir = Path(MARKDOWN_DIR)
    total_docs = len(changed_files)
    
    for idx, f_path in enumerate(changed_files, 1):
        logger.info("Converting [%d/%d] %s to Markdown...", idx, total_docs, f_path.name)
        try:
            if f_path.suffix.lower() == ".pdf":
                pdf_to_markdown(f_path, output_dir)
            elif f_path.suffix.lower() == ".docx":
                docx_to_markdown(f_path, output_dir)
        except Exception as e:
            logger.error("Failed to convert %s: %s", f_path.name, e)

# --- 2. Hierarchical Chunking Utilities ---
def merge_metadata(target: dict, source: dict, prepend: bool = False):
    for key, value in source.items():
        if key not in target:
            target[key] = value
        elif prepend:
            target[key] = f"{value} -> {target[key]}"
        else:
            target[key] = f"{target[key]} -> {value}"

def merge_small_parents(chunks: list, min_size: int) -> list:
    if not chunks:
        return []

    merged, current = [], None

    for chunk in chunks:
        if current is None:
            current = chunk
        else:
            current.page_content += "\n\n" + chunk.page_content
            merge_metadata(current.metadata, chunk.metadata)

        if len(current.page_content) >= min_size:
            merged.append(current)
            current = None

    if current:
        if merged:
            merged[-1].page_content += "\n\n" + current.page_content
            merge_metadata(merged[-1].metadata, current.metadata)
        else:
            merged.append(current)

    return merged

def split_large_parents(chunks: list, max_size: int, splitter) -> list:
    split_chunks = []

    # Domain-structure separators come from the active profile (e.g. legal
    # clause markers); generic text separators are used as the fallback.
    large_splitter = RecursiveCharacterTextSplitter(
        separators=PROFILE["chunk_separators"],
        chunk_size=max_size,
        chunk_overlap=splitter._chunk_overlap
    )

    for chunk in chunks:
        if len(chunk.page_content) <= max_size:
            split_chunks.append(chunk)
        else:
            split_chunks.extend(large_splitter.split_documents([chunk]))

    return split_chunks

def clean_small_chunks(chunks: list, min_size: int) -> list:
    cleaned = []

    for i, chunk in enumerate(chunks):
        if len(chunk.page_content) < min_size:
            if cleaned:
                cleaned[-1].page_content += "\n\n" + chunk.page_content
                merge_metadata(cleaned[-1].metadata, chunk.metadata)
            elif i < len(chunks) - 1:
                chunks[i + 1].page_content = chunk.page_content + "\n\n" + chunks[i + 1].page_content
                merge_metadata(chunks[i + 1].metadata, chunk.metadata, prepend=True)
            else:
                cleaned.append(chunk)
        else:
            cleaned.append(chunk)

    return cleaned

# --- 3. Vector Database Setup ---
def ensure_collection(collection_name: str):
    if not client.collection_exists(collection_name):
        dim = len(dense_embeddings.embed_query("test"))
        client.create_collection(
            collection_name=collection_name,
            vectors_config=qmodels.VectorParams(
                size=dim,
                distance=qmodels.Distance.COSINE
            ),
            sparse_vectors_config={
                "sparse": qmodels.SparseVectorParams()
            },
        )

# Initialize vector store instance
ensure_collection(CHILD_COLLECTION)

child_vector_store = QdrantVectorStore(
    client=client,
    collection_name=CHILD_COLLECTION,
    embedding=dense_embeddings,
    sparse_embedding=sparse_embeddings,
    retrieval_mode=RetrievalMode.HYBRID,
    sparse_vector_name="sparse"
)

# --- 4. Main Indexing Logic ---
def index_documents(changed_files: list):
    """Reads Markdown files, splits hierarchically, and indexes them."""
    if not changed_files:
        logger.info("No changed documents to index.")
        return
        
    logger.info("Starting document indexing for %d changed files...", len(changed_files))
    headers_to_split_on = [("#", "H1"), ("##", "H2"), ("###", "H3")]
    parent_splitter = MarkdownHeaderTextSplitter(headers_to_split_on=headers_to_split_on, strip_headers=False)
    
    # Child chunks are kept small for sharper embedding vectors
    child_splitter = RecursiveCharacterTextSplitter(
        separators=["\n\n", "\n", ". ", " ", ""],
        chunk_size=CHILD_CHUNK_SIZE,
        chunk_overlap=CHILD_CHUNK_OVERLAP
    )

    min_parent_size = MIN_PARENT_SIZE
    max_parent_size = MAX_PARENT_SIZE

    all_parent_pairs, all_child_chunks = [], []
    
    # Convert changed_files to markdown paths
    md_files = []
    for f_path in changed_files:
        md_path = (Path(MARKDOWN_DIR) / f_path.stem).with_suffix(".md")
        if md_path.exists():
            md_files.append(md_path)

    total_md_files = len(md_files)
    for idx, doc_path in enumerate(md_files, 1):
        logger.info("Processing markdown [%d/%d]: %s...", idx, total_md_files, doc_path.name)
        try:
            with open(doc_path, "r", encoding="utf-8") as f:
                md_text = f.read()
        except Exception as e:
            logger.error("Failed to read %s: %s", doc_path, e)
            continue

        parent_chunks = parent_splitter.split_text(md_text)
        merged_parents = merge_small_parents(parent_chunks, min_parent_size)
        split_parents = split_large_parents(merged_parents, max_parent_size, child_splitter)
        cleaned_parents = clean_small_chunks(split_parents, min_parent_size)

        # Determine correct source extension
        source_ext = ".docx" if os.path.exists(os.path.join(DOCS_DIR, f"{doc_path.stem}.docx")) else ".pdf"
        source_filename = f"{doc_path.stem}{source_ext}"

        # Delete old parent JSON files corresponding to this file from parent_store
        for item in os.listdir(PARENT_STORE_PATH):
            if item.startswith(f"{doc_path.stem}_parent_") and item.endswith(".json"):
                try:
                    os.remove(os.path.join(PARENT_STORE_PATH, item))
                except OSError as e:
                    logger.warning("Could not remove stale parent chunk %s: %s", item, e)

        for i, p_chunk in enumerate(cleaned_parents):
            parent_id = f"{doc_path.stem}_parent_{i}"
            p_chunk.metadata.update({"source": source_filename, "parent_id": parent_id})
            all_parent_pairs.append((parent_id, p_chunk))
            
            children = child_splitter.split_documents([p_chunk])
            all_child_chunks.extend(children)

    if not all_child_chunks:
        logger.info("No chunks generated.")
        return

    # Delete old vector chunks from Qdrant for these files before adding new ones
    for f_path in changed_files:
        source_filename = f_path.name
        logger.info("Deleting old Qdrant chunks for: %s", source_filename)
        try:
            client.delete(
                collection_name=CHILD_COLLECTION,
                points_selector=qmodels.Filter(
                    must=[
                        qmodels.FieldCondition(
                            key="metadata.source",
                            match=qmodels.MatchValue(value=source_filename)
                        )
                    ]
                )
            )
        except Exception as e:
            logger.error("Failed to delete old chunks from Qdrant: %s", e)

    # 4.1 Index child chunks into Qdrant in batches
    total_chunks = len(all_child_chunks)
    batch_size = INDEX_BATCH_SIZE
    logger.info("Indexing %d child chunks to Qdrant in batches of %d...", total_chunks, batch_size)
    try:
        import torch
        for idx in range(0, total_chunks, batch_size):
            batch = all_child_chunks[idx : idx + batch_size]
            child_vector_store.add_documents(batch)
            progress = min(idx + batch_size, total_chunks)
            percent = (progress / total_chunks) * 100
            logger.info("Indexed [%d/%d] child chunks (%.1f%%)", progress, total_chunks, percent)
            
            # Clear MPS cache to prevent memory fragmentation Out-Of-Memory errors
            if hasattr(torch.backends, "mps") and torch.backends.mps.is_available():
                torch.mps.empty_cache()
    except Exception as e:
        logger.error("Failed to index child chunks: %s", e)
        return

    # 4.2 Save parent chunks to local JSON files
    logger.info("Saving %d parent chunks to JSON...", len(all_parent_pairs))
    for parent_id, doc in all_parent_pairs:
        doc_dict = {"page_content": doc.page_content, "metadata": doc.metadata}
        filepath = os.path.join(PARENT_STORE_PATH, f"{parent_id}.json")
        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(doc_dict, f, ensure_ascii=False, indent=2)

    logger.info("Indexing completed successfully!")

if __name__ == "__main__":
    import argparse
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )
    parser = argparse.ArgumentParser(description="Index documents to Qdrant.")
    parser.add_argument("--clean", action="store_true", help="Clear existing index and markdown cache before indexing.")
    parser.add_argument("--overwrite", action="store_true", help="Force overwrite markdown cache files.")
    args = parser.parse_args()

    metadata_path = os.path.join(PARENT_STORE_PATH, "indexing_metadata.json")
    old_metadata = {}
    if os.path.exists(metadata_path) and not args.clean:
        try:
            with open(metadata_path, "r", encoding="utf-8") as f:
                old_metadata = json.load(f)
        except (OSError, json.JSONDecodeError) as e:
            logger.warning("Could not read index metadata %s, treating all files as changed: %s", metadata_path, e)
    old_hashes = old_metadata.get("file_hashes", {})

    if args.clean:
        logger.info("Cleaning existing database collection and cache...")
        if client.collection_exists(CHILD_COLLECTION):
            client.delete_collection(CHILD_COLLECTION)
        ensure_collection(CHILD_COLLECTION)

        # Clean parent_store
        for item in os.listdir(PARENT_STORE_PATH):
            try:
                os.remove(os.path.join(PARENT_STORE_PATH, item))
            except OSError as e:
                logger.warning("Could not remove %s: %s", item, e)

        # Clean markdown_docs
        for item in os.listdir(MARKDOWN_DIR):
            try:
                os.remove(os.path.join(MARKDOWN_DIR, item))
            except OSError as e:
                logger.warning("Could not remove %s: %s", item, e)

        old_hashes = {}

    # Gather all source files in docs/
    pdf_files = sorted(glob.glob(os.path.join(DOCS_DIR, "*.pdf")))
    docx_files = sorted(glob.glob(os.path.join(DOCS_DIR, "*.docx")))
    txt_files = sorted(glob.glob(os.path.join(DOCS_DIR, "*.txt")))
    md_files = sorted(glob.glob(os.path.join(DOCS_DIR, "*.md")))
    all_source_files = [Path(p) for p in pdf_files + docx_files + txt_files + md_files]

    # Calculate hashes and identify changed files
    current_hashes = {}
    changed_files = []
    for f_path in all_source_files:
        f_hash = calculate_md5(f_path)
        current_hashes[f_path.name] = f_hash
        md_path = (Path(MARKDOWN_DIR) / f_path.stem).with_suffix(".md")
        if args.clean or args.overwrite or f_path.name not in old_hashes or old_hashes[f_path.name] != f_hash or not md_path.exists():
            changed_files.append(f_path)
        else:
            logger.info("Skipping unchanged file: %s", f_path.name)

    if changed_files:
        import shutil
        # Copy direct text/md files to MARKDOWN_DIR
        for f_path in changed_files:
            if f_path.suffix.lower() in [".txt", ".md"]:
                md_path = (Path(MARKDOWN_DIR) / f_path.stem).with_suffix(".md")
                shutil.copy2(f_path, md_path)

        # Converts changed documents in 'docs/' to Markdown
        documents_to_markdowns(changed_files)
        index_documents(changed_files)
        
        # Save updated hashes metadata
        new_metadata = {
            "file_hashes": {**old_hashes, **{f.name: current_hashes[f.name] for f in changed_files}}
        }
        with open(metadata_path, "w", encoding="utf-8") as f:
            json.dump(new_metadata, f, ensure_ascii=False, indent=2)
    else:
        logger.info("All documents are up-to-date. Nothing to index.")
